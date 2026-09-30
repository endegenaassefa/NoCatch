from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageDraw
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader
from zipfile import ZipFile, ZIP_DEFLATED, ZipInfo
import io, json, hashlib
ROOT=Path(__file__).resolve().parent/'fixtures'
ROOT.mkdir(exist_ok=True)
image=Image.new('RGB',(900,300),'white'); draw=ImageDraw.Draw(image)
draw.text((30,80),'VISUAL ONLY: ORBITAL diagram shows 73 units',fill='black',font_size=34)
png=io.BytesIO(); image.save(png,format='PNG'); png.seek(0)
prs=Presentation()
a=prs.slides.add_slide(prs.slide_layouts[6])
a.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text='PHYSICAL ONE: Xenon calibration takes 17 minutes.'
a.notes_slide.notes_text_frame.text='PRIVATE_NOTES_SENTINEL: Xenon safety cutoff is 42 kelvin.'
table=a.shapes.add_table(2,2, Inches(1), Inches(3), Inches(6), Inches(1)).table
for r,row in enumerate([['Metric','Value'],['Xenon reserve','318 liters']]):
 for c,text in enumerate(row): table.cell(r,c).text=text
b=prs.slides.add_slide(prs.slide_layouts[6])
b.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text='LOGICAL FIRST: Zephyr launch code is MARIGOLD-629.'
b.shapes.add_picture(png, Inches(1), Inches(2),width=Inches(8))
b.notes_slide.notes_text_frame.text='FIRST_SLIDE_NOTE: Zephyr contingency takes 23 hours.'
c=prs.slides.add_slide(prs.slide_layouts[6])
c.shapes.add_picture(png, Inches(1), Inches(1),width=Inches(8))
# Physical slide2 comes first in presentation.xml, then slide1, then image-only slide3.
ids=prs.slides._sldIdLst; ids.insert(0,ids[1])
buf=io.BytesIO(); prs.save(buf)
# Normalize ZIP container timestamps to make generated fixture hashes repeatable.
def zipwrite(name, entries):
 with ZipFile(ROOT/name,'w',compression=ZIP_DEFLATED) as z:
  for key,data in sorted(entries.items()):
   info=ZipInfo(key,(2020,1,1,0,0,0));info.compress_type=ZIP_DEFLATED;z.writestr(info,data)
with ZipFile(io.BytesIO(buf.getvalue())) as z: entries={n:z.read(n) for n in z.namelist()}
zipwrite('relationship-order-notes-visual.pptx',entries)
zipwrite('inflated-text.pptx',{**entries,'ppt/slides/slide1.xml':entries['ppt/slides/slide1.xml'].replace(b'PHYSICAL ONE:', b'Z'*(5*1024*1024)+b' PHYSICAL ONE:')})
def pdf(name,pages):
 c=Canvas(str(ROOT/name),invariant=1)
 for lines,visual in pages:
  y=750
  for line in lines: c.drawString(50,y,line); y-=25
  if visual: c.drawImage(ImageReader(image),50,350,width=450,height=150)
  c.showPage()
 c.save()
pdf('three-pages.pdf',[(['PDF_ONE_SENTINEL: Quasar reserve is 961 liters.'],False),(['PDF_TWO_SENTINEL: Aurora deadline is Friday at noon.'],True),([],True)])
pdf('251-pages.pdf',[([f'Limit fixture page {n+1}.'],False) for n in range(251)])
for n in range(11): pdf(f'unique-{n+1}.pdf',[([f'UNIQUE_{n+1}: specimen identifier {n+1}.'],False)])
(ROOT/'corrupt.pptx').write_bytes(b'PK\x03\x04broken archive')
(ROOT/'corrupt.pdf').write_bytes(b'%PDF-1.7\ninvalid object graph')
(ROOT/'legacy.ppt').write_bytes(bytes.fromhex('D0CF11E0A1B11AE1')+b'unsupported legacy container')
manifest={p.name:{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(ROOT.iterdir()) if p.is_file() and p.name!='manifest.json'}
(ROOT/'manifest.json').write_text(json.dumps({'expected':{'pptx':{'1':'LOGICAL FIRST / MARIGOLD-629 / FIRST_SLIDE_NOTE','2':'PHYSICAL ONE / 17 minutes / 318 liters / PRIVATE_NOTES_SENTINEL','3':'image only; partial/unavailable coverage'},'pdf':{'1':'PDF_ONE_SENTINEL / 961 liters','2':'PDF_TWO_SENTINEL plus image gap','3':'image-only coverage gap'}},'files':manifest},indent=2)+'\n')
print(json.dumps({'fixtures':len(manifest),'root':str(ROOT)}))
