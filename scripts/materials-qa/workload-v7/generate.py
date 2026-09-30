from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from PIL import Image,ImageDraw
import io,json,hashlib
root=Path(__file__).resolve().parent/'results'/'decks';root.mkdir(parents=True,exist_ok=True)
picture=Image.new('RGB',(600,200),'#edf0f6');d=ImageDraw.Draw(picture);d.rectangle((40,50,230,150),fill='#5a76ae');d.rectangle((350,20,570,150),fill='#75ac8c');d.text((50,165),'Visual-only capacity chart',fill='black');png=io.BytesIO();picture.save(png,format='PNG')
for deck in range(1,11):
 prs=Presentation()
 for page in range(1,61):
  slide=prs.slides.add_slide(prs.slide_layouts[6])
  if page%15==0:
   png.seek(0);slide.shapes.add_picture(png, Inches(1), Inches(1),width=Inches(8));continue
  title=f'Workshop {deck}: systems planning, section {page}'
  slide.shapes.add_textbox(Inches(.6), Inches(.4), Inches(8.5), Inches(.7)).text=title
  body=(f'Workshop {deck}, section {page} examines capacity planning for a regional service. '
   'The planning team records demand estimates, owner responsibilities, verification evidence and delivery constraints. '
   'Each proposal must distinguish measured behavior from assumptions and document what would invalidate the estimate. '
   'The weekly review compares resource availability with outstanding tasks, tests failure recovery, and records follow-up decisions. '
   'The schedule includes preparation, independent review, a controlled pilot and an explicit checkpoint before wider adoption. '
   'Keep these observations attached to their source section so a reviewer can trace each answer to its context.')
  if deck==10 and page==55:body+=' HELIOSDECK10SLIDE55: the reserve threshold is exactly 9473 liters and the approval codename is COBALT-MEADOW.'
  box=slide.shapes.add_textbox(Inches(.6), Inches(1.3), Inches(8.5), Inches(3.4));box.text=body
  for paragraph in box.text_frame.paragraphs:paragraph.font.size=Pt(17)
  slide.notes_slide.notes_text_frame.text=(f'Speaker notes for workshop {deck}, slide {page}: ask the group to describe one operational risk and one observable acceptance check. '
    f'The section-specific review ticket is WK{deck:02d}-S{page:02d}. Clarify any uncertain assumption before citing this section as an established result.')
  if page%10==0:
   tab=slide.shapes.add_table(3,2,Inches(.8), Inches(5), Inches(7), Inches(1)).table
   for r,row in enumerate([['Measure','Recorded value'],['Planning cycle',str(deck*100+page)],['Review owners',str(deck+page)]]):
    for c,value in enumerate(row):tab.cell(r,c).text=value
  elif page%8==0:
   png.seek(0);slide.shapes.add_picture(png,Inches(1), Inches(5),width=Inches(5))
 prs.save(root/f'workshop-{deck:02d}.pptx')
files={p.name:{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(root.glob('*.pptx'))}
(root/'manifest.json').write_text(json.dumps({'decks':10,'slidesPerDeck':60,'totalSlides':600,'fact':{'file':'workshop-10.pptx','page':55,'anchor':'HELIOSDECK10SLIDE55','answer':'9473 liters','codename':'COBALT-MEADOW'},'files':files},indent=2)+'\n');print(json.dumps({'files':10,'slides':600,'sourceBytes':sum(x['bytes'] for x in files.values())}))
